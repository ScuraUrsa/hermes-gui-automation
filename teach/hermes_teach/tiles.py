"""Built-in workflow tiles ("uses: <name>"): the ready-made blocks between recorded procedures.

Every tile is a function ``(step_config, ctx, runtime) -> outputs`` where
``step_config`` is the step's own fields already rendered with the run
context, ``ctx`` the run context and ``runtime`` gives access to the channel,
the LLM, the dry-run flag and a logger. Outputs become ``steps.<id>.<name>``.

Tiles: ``shell``, ``http``, ``llm``, ``ask``, ``notify``, ``set``,
``table.append`` (CSV / XLSX row), ``wait``. ``procedure`` and ``workflow``
are handled by the workflow runner itself because they need drivers.
"""

from __future__ import annotations

import csv
import json
import os
import re
import shlex
import subprocess
import time
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from hermes_teach.channels import Channel, is_yes
from hermes_teach.procedure import ProcedureError


class TileError(RuntimeError):
    pass


@dataclass
class Runtime:
    channel: Channel
    dry_run: bool = False
    log: Callable[[str], None] = lambda _m: None
    llm: Any = None                  # hermes_teach.llm.LLM, created lazily
    llm_factory: Callable[[], Any] | None = None
    ask_timeout: float = 3600

    def get_llm(self) -> Any:
        if self.llm is None:
            if self.llm_factory is None:
                from hermes_teach.llm import from_env
                self.llm_factory = from_env
            self.llm = self.llm_factory()
        return self.llm


def parse_number(value: Any) -> float | None:
    """'1 234,56 zł' -> 1234.56; '1,234.56' -> 1234.56; None if no number."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value)
    m = re.search(r"-?\d[\d  .,]*", s)
    if not m:
        return None
    num = m.group(0).replace(" ", "").replace(" ", "")
    if "," in num and "." in num:
        num = num.replace(",", "") if num.rfind(".") > num.rfind(",") else num.replace(".", "").replace(",", ".")
    elif "," in num:
        num = num.replace(",", ".") if num.count(",") == 1 and len(num.split(",")[1]) <= 2 else num.replace(",", "")
    elif num.count(".") > 1:
        num = num.replace(".", "")
    try:
        return float(num)
    except ValueError:
        return None


# --- tiles -------------------------------------------------------------------
def tile_shell(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    cmd = cfg.get("run")
    if not cmd:
        raise ProcedureError("shell tile needs 'run'")
    if rt.dry_run:
        rt.log(f"would run: {cmd}")
        return {"code": 0, "stdout": "", "stderr": "", "json": None}
    env = {**os.environ, **{str(k): str(v) for k, v in (cfg.get("env") or {}).items()}}
    args = str(cmd) if cfg.get("shell", True) else shlex.split(str(cmd), posix=os.name != "nt")
    try:
        res = subprocess.run(args, shell=bool(cfg.get("shell", True)), capture_output=True, text=True,
                             timeout=float(cfg.get("timeout", 600)), env=env, cwd=cfg.get("cwd") or None)
    except subprocess.TimeoutExpired:
        raise TileError(f"command timed out: {cmd}") from None
    out = res.stdout.strip()
    parsed: Any = None
    if out[:1] in ("{", "["):
        try:
            parsed = json.loads(out)
        except json.JSONDecodeError:
            parsed = None
    if res.returncode != 0 and not cfg.get("allow_failure"):
        raise TileError(f"command exited {res.returncode}: {res.stderr.strip()[:300] or out[:300]}")
    return {"code": res.returncode, "stdout": out, "stderr": res.stderr.strip(), "json": parsed}


def tile_http(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    url = cfg.get("url")
    if not url:
        raise ProcedureError("http tile needs 'url'")
    method = str(cfg.get("method", "POST" if ("json" in cfg or "body" in cfg) else "GET")).upper()
    headers = {str(k): str(v) for k, v in (cfg.get("headers") or {}).items()}
    data: bytes | None = None
    if "json" in cfg:
        data = json.dumps(cfg["json"], ensure_ascii=False).encode()
        headers.setdefault("Content-Type", "application/json")
    elif "body" in cfg:
        data = str(cfg["body"]).encode()
    if rt.dry_run:
        rt.log(f"would {method} {url}")
        return {"status": 0, "body": "", "json": None}
    req = urllib.request.Request(str(url), data=data, method=method, headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=float(cfg.get("timeout", 60))) as r:
            body = r.read().decode("utf-8", "replace")
            status = r.status
    except urllib.error.HTTPError as e:
        body, status = e.read().decode("utf-8", "replace"), e.code
        if not cfg.get("allow_failure"):
            raise TileError(f"{method} {url}: HTTP {status} {body[:200]}") from None
    except urllib.error.URLError as e:
        raise TileError(f"{method} {url}: {e.reason}") from None
    parsed: Any = None
    try:
        parsed = json.loads(body) if body.strip()[:1] in ("{", "[") else None
    except json.JSONDecodeError:
        parsed = None
    return {"status": status, "body": body, "json": parsed}


def tile_llm(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    prompt = cfg.get("prompt")
    if not prompt:
        raise ProcedureError("llm tile needs 'prompt'")
    system = str(cfg.get("system") or "You are a precise assistant inside an automation workflow. "
                 "Answer with the requested content only, no preamble.")
    if cfg.get("json"):
        system += " Answer with a single JSON object only."
    if rt.dry_run:
        rt.log(f"would ask the LLM: {str(prompt)[:80]}")
        return {"text": "", "json": {}}
    text = rt.get_llm().complete(system, str(prompt))
    parsed: Any = None
    if cfg.get("json"):
        m = re.search(r"\{.*\}|\[.*\]", text, re.S)
        try:
            parsed = json.loads(m.group(0)) if m else None
        except json.JSONDecodeError:
            parsed = None
        if parsed is None:
            raise TileError("LLM did not return JSON")
    return {"text": text.strip(), "json": parsed}


def tile_ask(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    question = cfg.get("question")
    if not question:
        raise ProcedureError("ask tile needs 'question'")
    options = [str(o) for o in (cfg.get("options") or [])]
    answer = rt.channel.ask(str(question), options or None, timeout=float(cfg.get("timeout", rt.ask_timeout)))
    if answer is None and cfg.get("default") is not None:
        answer = str(cfg["default"])
    return {"answer": answer, "yes": is_yes(answer) if not options else (answer == options[0]),
            "answered": answer is not None}


def tile_notify(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    text = cfg.get("text")
    if text is None:
        raise ProcedureError("notify tile needs 'text'")
    rt.channel.notify(str(text))
    return {"text": str(text)}


def tile_set(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    values = cfg.get("values")
    if not isinstance(values, dict):
        raise ProcedureError("set tile needs a 'values' mapping")
    return dict(values)


def tile_wait(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    secs = float(cfg.get("seconds", 0))
    if not rt.dry_run and secs > 0:
        time.sleep(secs)
    return {"seconds": secs}


def tile_table_append(cfg: dict[str, Any], ctx: dict[str, Any], rt: Runtime) -> dict[str, Any]:
    """Append one row to a CSV (stdlib) or XLSX (openpyxl) file; creates it with a header row."""
    file = cfg.get("file")
    row = cfg.get("row")
    if not file or not isinstance(row, dict):
        raise ProcedureError("table.append needs 'file' and a 'row' mapping")
    path = Path(str(file)).expanduser()
    if rt.dry_run:
        rt.log(f"would append to {path}: {row}")
        return {"file": str(path), "row": dict(row), "rows": 0}
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise TileError("pip install openpyxl to write .xlsx files (or use a .csv file)") from None
        sheet_name = cfg.get("sheet")
        if path.exists():
            wb = openpyxl.load_workbook(path)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else (
                wb.create_sheet(sheet_name) if sheet_name else wb.active)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            if sheet_name:
                ws.title = sheet_name
        header = [c.value for c in ws[1]] if ws.max_row >= 1 and any(c.value for c in ws[1]) else []
        if not header:
            header = list(row)
            ws.append(header)
        for k in row:
            if k not in header:
                header.append(k)
                ws.cell(row=1, column=len(header), value=k)
        ws.append([row.get(h) for h in header])
        wb.save(path)
        return {"file": str(path), "row": dict(row), "rows": ws.max_row - 1}
    exists = path.exists() and path.stat().st_size > 0
    header: list[str] = []
    if exists:
        with open(path, newline="", encoding="utf-8") as f:
            header = next(csv.reader(f), [])
    if not header:
        header = list(row)
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if not exists:
            w.writerow(header)
        w.writerow([row.get(h, "") for h in header])
    with open(path, newline="", encoding="utf-8") as f:
        n = sum(1 for _ in f) - 1
    return {"file": str(path), "row": dict(row), "rows": n}


TILES: dict[str, Callable[[dict[str, Any], dict[str, Any], Runtime], dict[str, Any]]] = {
    "shell": tile_shell, "http": tile_http, "llm": tile_llm, "ask": tile_ask, "notify": tile_notify,
    "set": tile_set, "wait": tile_wait, "table.append": tile_table_append,
}

TILE_DOCS: dict[str, str] = {
    "procedure": "procedure: <name|path>, with: {param: value}  -> status, ok, run_id, executed",
    "workflow": "workflow: <name|path>, with: {input: value}     -> status, ok + the workflow's outputs",
    "shell": "run: <command>, env?, cwd?, timeout?, allow_failure? -> code, stdout, stderr, json",
    "http": "url, method?, json?|body?, headers?, allow_failure? -> status, body, json",
    "llm": "prompt, system?, json? -> text, json",
    "ask": "question, options?: [..], default?, timeout? -> answer, yes, answered   (goes to the user's channel)",
    "notify": "text -> text   (message to the user's channel)",
    "set": "values: {name: value-or-template} -> each value",
    "wait": "seconds -> seconds",
    "table.append": "file: x.csv|x.xlsx, sheet?, row: {column: value} -> file, rows",
}
